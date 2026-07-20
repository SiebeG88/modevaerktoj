"""Word-eksport (.docx) af referat + transkription.

Ren, GUI-fri renderer bygget på python-docx (ingen ekstern afhængighed som
pandoc — virker derfor også på Windows uden ekstra installation). Konverterer
referatets lette markdown til Word-elementer og skriver et metadata-hoved med
mødenavn, dato, mødetype og deltagere."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

# Inline-tokens: **fed** og *kursiv*. Capturing-group så re.split beholder dem.
_INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")
_HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)$")
_BULLET_RE = re.compile(r"^[-*]\s+(.*)$")
_NUMBER_RE = re.compile(r"^\d+\.\s+(.*)$")


def _add_inline(paragraph, text: str) -> None:
    """Tilføj tekst til et afsnit med **fed**/*kursiv* som separate runs."""
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if len(part) >= 4 and part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif len(part) >= 2 and part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _render_markdown(doc, body: str) -> None:
    """Render referatets lette markdown til Word-elementer. Ukendte linjer
    bliver almindelige afsnit; blanklinjer springes over (ingen tomme afsnit)."""
    for raw in body.splitlines():
        line = raw.strip()
        if not line:
            continue
        m = _HEADING_RE.match(line)
        if m:
            doc.add_heading(m.group(2), level=len(m.group(1)))
            continue
        m = _BULLET_RE.match(line)
        if m:
            _add_inline(doc.add_paragraph(style="List Bullet"), m.group(1))
            continue
        m = _NUMBER_RE.match(line)
        if m:
            _add_inline(doc.add_paragraph(style="List Number"), m.group(1))
            continue
        if line.startswith(">"):
            _add_inline(doc.add_paragraph(style="Intense Quote"), line[1:].strip())
            continue
        _add_inline(doc.add_paragraph(), line)


def _add_meta_line(doc, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(value)


def write_meeting_docx(
    path: Path,
    *,
    title: str,
    date: str,
    meeting_type_name: str | None,
    attendees: list[str],
    body: str,
    is_transcript: bool,
) -> Path:
    """Skriv ét mødedokument (.docx) med metadata-hoved + indhold.

    Referatet (is_transcript=False) renderes fra markdown; transkriptionen
    (is_transcript=True) skrives som ét afsnit pr. ikke-tom linje, så
    talerlinjer bevares. Returnerer path."""
    doc = Document()
    doc.add_heading(title, level=0)  # Word "Title"-stil
    _add_meta_line(doc, "Dato", date)
    if meeting_type_name:
        _add_meta_line(doc, "Mødetype", meeting_type_name)
    if attendees:
        _add_meta_line(doc, "Deltagere", ", ".join(attendees))
    doc.add_paragraph()  # luft før indhold

    if is_transcript:
        for raw in body.splitlines():
            line = raw.strip()
            if line:
                doc.add_paragraph(line)
    else:
        _render_markdown(doc, body)

    doc.save(str(path))
    return path
